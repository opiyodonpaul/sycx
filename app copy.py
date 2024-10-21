from flask import Flask, request, jsonify
from flask_socketio import SocketIO, emit
from summarie import summarize_documents, enrich_summary_with_visuals, convert_summary_to_pdf
from model import get_model
from dotenv import load_dotenv
import os
import base64
import io
from concurrent.futures import ThreadPoolExecutor
import logging
import json
import traceback

load_dotenv()

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 1024 * 1024 * 1024
socketio = SocketIO(app, cors_allowed_origins="*", max_http_buffer_size=1024 * 1024 * 1024)
summarization_model = get_model()

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

executor = ThreadPoolExecutor(max_workers=10)

def extract_text_from_document(content, file_type):
    try:
        if isinstance(content, str):
            try:
                content = base64.b64decode(content)
            except:
                content = content.encode('utf-8')
        
        if file_type == 'pdf':
            from pdfminer.high_level import extract_text
            return extract_text(io.BytesIO(content))
        elif file_type in ['doc', 'docx']:
            from docx import Document
            doc = Document(io.BytesIO(content))
            return "\n".join([para.text for para in doc.paragraphs])
        elif file_type in ['xls', 'xlsx']:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content))
            text = ""
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    text += " | ".join([str(cell) for cell in row if cell]) + "\n"
            return text
        elif file_type in ['ppt', 'pptx']:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(content))
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        text += shape.text + "\n"
            return text
        elif file_type in ['png', 'jpg', 'jpeg']:
            import pytesseract
            from PIL import Image
            image = Image.open(io.BytesIO(content))
            return pytesseract.image_to_string(image)
        else:
            return content.decode('utf-8', errors='ignore')
    except Exception as e:
        logging.error(f"Error extracting text from {file_type} file: {str(e)}")
        logging.error(traceback.format_exc())
        raise Exception(f"Error extracting text from {file_type} file: {str(e)}")

def progress_callback(progress):
    socketio.emit('summarization_progress', {'type': 'progress', 'progress': progress})

@app.route('/summarize', methods=['POST'])
def summarize():
    try:
        data = request.get_json()
        if not data:
            return jsonify({'error': 'No JSON data received'}), 400

        merge_summaries = data.get('merge_summaries', False)
        summary_depth = float(data.get('summary_depth', 0.3))
        language = data.get('language', 'en')
        documents_data = data.get('documents', [])

        if not documents_data:
            return jsonify({'error': 'No documents provided'}), 400

        documents = []
        for doc in documents_data:
            try:
                text = extract_text_from_document(doc['content'], doc['type'])
                documents.append({
                    'name': doc['name'],
                    'content': text,
                    'type': doc['type']
                })
            except Exception as e:
                logging.error(f"Error processing document {doc['name']}: {str(e)}")
                logging.error(traceback.format_exc())
                return jsonify({'error': f"Error processing document {doc['name']}: {str(e)}"}), 400

        summaries = summarize_documents(
            summarization_model, 
            documents, 
            merge_summaries, 
            summary_depth, 
            language, 
            progress_callback
        )
        
        enriched_summaries = [enrich_summary_with_visuals(summary) for summary in summaries]
        
        pdf_summaries = []
        for summary in enriched_summaries:
            try:
                pdf_content = convert_summary_to_pdf(summary['content'])
                pdf_base64 = base64.b64encode(pdf_content).decode('utf-8')
                pdf_summaries.append({
                    'title': summary['title'],
                    'content': pdf_base64
                })
            except Exception as e:
                logging.error(f"Error converting summary to PDF: {str(e)}")
                pdf_summaries.append({
                    'title': summary['title'],
                    'content': summary['content']
                })
        
        return jsonify({'summaries': pdf_summaries}), 200

    except Exception as e:
        logging.error(f"Error in summarize route: {str(e)}")
        logging.error(traceback.format_exc())
        return jsonify({'error': str(e)}), 500

@socketio.on('connect')
def handle_connect():
    logging.info('Client connected')
    emit('connected', {'data': 'Connected'})

@socketio.on('disconnect')
def handle_disconnect():
    logging.info('Client disconnected')

@app.route('/feedback', methods=['POST'])
def feedback():
    data = request.json
    summary_id = data.get('summary_id')
    user_id = data.get('user_id')
    feedback_text = data.get('feedback')

    if not all([summary_id, user_id, feedback_text]):
        return jsonify({'error': 'Missing required data'}), 400

    try:
        # Here you would typically update the model based on the feedback
        return jsonify({'message': 'Feedback received successfully'}), 200
    except Exception as e:
        logging.error(f"Error in feedback route: {str(e)}")
        logging.error(traceback.format_exc())
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    socketio.run(app, debug=False, host='0.0.0.0', port=int(os.environ.get('PORT', 5000)))